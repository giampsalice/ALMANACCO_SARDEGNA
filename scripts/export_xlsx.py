#!/usr/bin/env python3
"""Convert the Almanacco workbook into the JSON consumed by the gallery."""

import json
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "source" / "Archivio_Almanacco_Sardegna_WordPress.xlsx"
OUTPUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data" / "almanacco.json"
CSV_FALLBACK = ROOT / "data" / "numeri.csv"


def text(value):
    return "" if value is None else str(value).strip()


wb = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = wb["Numeri"]
headers = [text(cell.value) for cell in sheet[4]]
rows = []
for values in sheet.iter_rows(min_row=5, values_only=True):
    if not values[0]:
        continue
    rows.append(dict(zip(headers, values)))

indices = {}
for sheet_name in wb.sheetnames:
    match = re.fullmatch(r"Indice (\d{4}(?:-\d{2})?)", sheet_name)
    if not match:
        continue
    ws = wb[sheet_name]
    contents = []
    issue_id = ""
    for values in ws.iter_rows(min_row=5, values_only=True):
        if not values[0]:
            continue
        issue_id = text(values[0])
        contents.append({
            "order": values[1],
            "section": text(values[2]),
            "author": text(values[3]),
            "title": text(values[4]),
            "pageStart": values[5],
            "pageEnd": values[6],
            "pages": text(values[7]),
            "notes": text(values[8]),
            "visible": text(values[9]).lower() != "no",
        })
    if issue_id:
        indices[issue_id] = contents

issues = []
for row in rows:
    issue_id = text(row.get("ID numero"))
    year = int(row.get("Anno"))
    title = text(row.get("Titolo"))
    label_match = re.search(r"\b(19|20)\d{2}(?:[/–-](?:19|20)?\d{2})?", title)
    year_label = label_match.group(0).replace("–", "/").replace("-", "/") if label_match else str(year)
    issues.append({
        "id": issue_id,
        "title": title,
        "year": year,
        "yearLabel": year_label,
        "decade": (year // 10) * 10,
        "volume": row.get("Volume numero"),
        "place": text(row.get("Luogo")),
        "publisher": text(row.get("Editore")),
        "lastPage": row.get("Ultima pagina indicizzata"),
        "slug": text(row.get("Slug")),
        "pdfUrl": text(row.get("URL download numero")),
        "coverUrl": text(row.get("URL immagine copertina")),
        "coverAlt": text(row.get("Testo alternativo copertina")) or f"Copertina {title}",
        "featured": text(row.get("In evidenza")).lower() != "no",
        "sortOrder": row.get("Ordinamento") or year,
        "sourceFile": text(row.get("File sorgente")),
        "notes": text(row.get("Note")),
        "contents": [item for item in indices.get(issue_id, []) if item["visible"]],
    })

issues.sort(key=lambda item: (item["sortOrder"], item["volume"] or 0))
payload = {
    "title": "Almanacco della Sardegna",
    "description": "Archivio digitale dei numeri dell’Almanacco della Sardegna",
    "generatedFrom": SOURCE.name,
    "count": len(issues),
    "issues": issues,
}
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
with CSV_FALLBACK.open("w", encoding="utf-8-sig", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        writer.writerow({header: row.get(header, "") for header in headers})
print(f"Exported {len(issues)} issues and {sum(len(x['contents']) for x in issues)} contents to {OUTPUT}")
